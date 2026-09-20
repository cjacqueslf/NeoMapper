"""Create localized NEOMapper manuals from the current Portuguese layout."""
from __future__ import annotations

import argparse
import json
import re
import urllib.parse
import urllib.request
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from docx import Document

from neomapper.shared.version import APP_VERSION


ROOT = Path(__file__).resolve().parents[1]
CACHE_PATH = ROOT / "docs" / "manual_translations.json"
LANGUAGES = {
    "en": ("English", "NEOMapper_User_Manual", "Page"),
    "es": ("Spanish", "Manual_del_Usuario_NEOMapper", "Página"),
}
PROTECTED = ("NEOMapper", "JPL Horizons", "Minor Planet Center", "MPES", "AstroNEOS", "SONEAR")
REVIEWED_TRANSLATIONS = {
    "en": {
        "Aba Info Obs": "Observing Information tab",
        "Info Obs": "Observing Information",
        "Passe para Info Obs, gere o gráfico de altitude e, se desejar, o Resumo Orbital.": "Go to Observing Information, generate the altitude plot and, if desired, the Orbital Summary.",
        "Uma busca bem-sucedida mostra o nome resolvido, a distância ao Sol, a distância à Terra e a magnitude. Ela também habilita Plotar Sistema Solar e Continuar para Info Obs. Alterar o objeto ou a data invalida o resultado anterior e exige uma nova busca.": "A successful search shows the resolved name, distance to the Sun, distance to Earth, and magnitude. It also enables Plot Solar System and Continue to Observing Information. Changing the object or date invalidates the previous result and requires a new search.",
        "Info Obs → limites → Plotar gráfico de altitude": "Observing Information → limits → Plot altitude graph",
        "Info Obs → Resumo orbital": "Observing Information → Orbital Summary",
    },
    "es": {
        "Aba Info Obs": "Pestaña Información de observación",
        "Info Obs": "Información de observación",
        "Passe para Info Obs, gere o gráfico de altitude e, se desejar, o Resumo Orbital.": "Vaya a Información de observación, genere el gráfico de altitud y, si lo desea, el Resumen orbital.",
        "Uma busca bem-sucedida mostra o nome resolvido, a distância ao Sol, a distância à Terra e a magnitude. Ela também habilita Plotar Sistema Solar e Continuar para Info Obs. Alterar o objeto ou a data invalida o resultado anterior e exige uma nova busca.": "Una búsqueda correcta muestra el nombre resuelto, la distancia al Sol, la distancia a la Tierra y la magnitud. También habilita Trazar el sistema solar y Continuar a Información de observación. Cambiar el objeto o la fecha invalida el resultado anterior y requiere una nueva búsqueda.",
        "Info Obs → limites → Plotar gráfico de altitude": "Información de observación → límites → Trazar gráfico de altitud",
        "Info Obs → Resumo orbital": "Información de observación → Resumen orbital",
        "Player": "Reproductor",
    },
}


def online_translation(text: str, language: str) -> str:
    if not text.strip() or text.startswith("%LOCALAPPDATA%"):
        return text.replace("<objeto>", "<object>" if language == "en" else "<objeto>").replace(
            "<data>", "<date>" if language == "en" else "<fecha>"
        )
    masked = text
    tokens: dict[str, str] = {}
    for index, value in enumerate(PROTECTED):
        token = f"ZXQ{index}QXZ"
        if value in masked:
            masked = masked.replace(value, token)
            tokens[token] = value
    query = urllib.parse.urlencode({"client": "gtx", "sl": "pt", "tl": language, "dt": "t", "q": masked})
    with urllib.request.urlopen(f"https://translate.googleapis.com/translate_a/single?{query}", timeout=30) as response:
        payload = json.load(response)
    result = "".join(part[0] for part in payload[0] if part[0])
    for token, value in tokens.items():
        result = result.replace(token, value).replace(token.lower(), value)
    return result


def collect_text(document: Document) -> list[str]:
    values = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        values.extend(cell.text for row in table.rows for cell in row.cells if cell.text.strip())
    return list(dict.fromkeys(values))


def load_translations(texts: list[str], language: str) -> dict[str, str]:
    cache = json.loads(CACHE_PATH.read_text(encoding="utf-8")) if CACHE_PATH.exists() else {}
    language_cache = cache.setdefault(language, {})
    language_cache.update(REVIEWED_TRANSLATIONS.get(language, {}))
    missing = [text for text in texts if text not in language_cache]
    if missing:
        with ThreadPoolExecutor(max_workers=6) as executor:
            translated = list(executor.map(lambda value: online_translation(value, language), missing))
        language_cache.update(zip(missing, translated))
        CACHE_PATH.write_text(json.dumps(cache, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return language_cache


def replace_paragraph(paragraph, translated: str) -> None:
    was_bold = bool(paragraph.runs) and all(run.bold for run in paragraph.runs)
    was_italic = bool(paragraph.runs) and all(run.italic for run in paragraph.runs)
    for run in paragraph.runs:
        run._element.getparent().remove(run._element)
    run = paragraph.add_run(translated)
    run.bold = was_bold
    run.italic = was_italic


def translate_document(source: Path, language: str, output: Path) -> None:
    document = Document(source)
    translations = load_translations(collect_text(document), language)
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            replace_paragraph(paragraph, translations[paragraph.text])
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if paragraph.text.strip():
                        replace_paragraph(paragraph, translations[paragraph.text])
    page_word = LANGUAGES[language][2]
    for section in document.sections:
        for paragraph in section.footer.paragraphs:
            for run in paragraph.runs:
                if run.text.strip().startswith("Página"):
                    run.text = re.sub(r"Página", page_word, run.text, count=1)
    document.core_properties.title = translations.get("Manual do Usuário do NEOMapper", document.core_properties.title)
    document.core_properties.subject = translations.get(
        f"Operação do NEOMapper versão {APP_VERSION}", f"NEOMapper {APP_VERSION} — {LANGUAGES[language][0]}"
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("language", choices=LANGUAGES)
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    translate_document(args.source, args.language, args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
