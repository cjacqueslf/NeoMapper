"""Build the current NEOMapper user manual as a reproducible DOCX."""
from __future__ import annotations

from pathlib import Path
import os
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from neomapper.shared.version import APP_VERSION


ROOT = Path(__file__).resolve().parents[1]
ASSETS = Path(os.environ.get("NEOMAPPER_MANUAL_ASSETS", ROOT / "docs" / "manual_assets"))
OUTPUT = Path(os.environ.get("NEOMAPPER_MANUAL_OUTPUT", ROOT / "docs" / f"Manual_do_Usuario_NEOMapper_{APP_VERSION}.docx"))
BLUE = "174A7E"
PALE_BLUE = "EEF5FB"
GRAY = "D9D9D9"


def shade(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    node = props.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        props.append(node)
    node.set(qn("w:fill"), fill)


def cell_margins(cell, value: int = 110) -> None:
    props = cell._tc.get_or_add_tcPr()
    margins = props.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        props.append(margins)
    for edge in ("top", "start", "bottom", "end"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)


def borders(table) -> None:
    props = table._tbl.tblPr
    grid = props.first_child_found_in("w:tblBorders")
    if grid is None:
        grid = OxmlElement("w:tblBorders")
        props.append(grid)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), GRAY)
        grid.append(node)


def heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True


def paragraph(document: Document, text: str, *, bold_lead: str | None = None) -> None:
    value = document.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        value.add_run(bold_lead).bold = True
        value.add_run(text[len(bold_lead):])
    else:
        value.add_run(text)


def bullet(document: Document, text: str) -> None:
    document.add_paragraph(text, style="List Bullet")


def numbered(document: Document, text: str) -> None:
    document.add_paragraph(text, style="List Number")


def table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    result = document.add_table(rows=1, cols=len(headers))
    result.autofit = False
    borders(result)
    header = result.rows[0]
    header._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for cell, text, width in zip(header.cells, headers, widths):
        cell.width = Inches(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(cell, BLUE)
        cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(9)
    for row_index, values in enumerate(rows):
        cells = result.add_row().cells
        for cell, text, width in zip(cells, values, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_margins(cell)
            if row_index % 2:
                shade(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.add_run(text).font.size = Pt(9)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def figure(document: Document, filename: str, caption: str, width: float = 6.8) -> None:
    path = ASSETS / filename
    if not path.is_file():
        raise FileNotFoundError(path)
    holder = document.add_paragraph()
    holder.alignment = WD_ALIGN_PARAGRAPH.CENTER
    holder.paragraph_format.keep_with_next = True
    shape = holder.add_run().add_picture(str(path), width=Inches(width))
    doc_pr = shape._inline.docPr
    doc_pr.set("descr", caption)
    cap = document.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(3)
    cap.paragraph_format.space_after = Pt(8)
    cap.runs[0].italic = True
    cap.runs[0].font.size = Pt(8)


def page_break(document: Document) -> None:
    document.add_page_break()


def remove_paragraph_borders(value) -> None:
    props = value._p.get_or_add_pPr()
    paragraph_borders = props.find(qn("w:pBdr"))
    if paragraph_borders is None:
        paragraph_borders = OxmlElement("w:pBdr")
        props.append(paragraph_borders)
    for edge in ("top", "left", "bottom", "right", "between", "bar"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        paragraph_borders.append(node)


def footer_page_number(section) -> None:
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run("Página ").font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    p._p.append(field)


def configure(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(.62)
    section.bottom_margin = Inches(.62)
    section.left_margin = Inches(.68)
    section.right_margin = Inches(.68)
    footer_page_number(section)
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.08
    for name, size in (("Title", 28), ("Subtitle", 14), ("Heading 1", 19), ("Heading 2", 13)):
        style = styles[name]
        style.font.name = "Aptos Display" if name != "Normal" else "Aptos"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(6)


def build_manual() -> None:
    document = Document()
    configure(document)

    # Cover
    document.add_paragraph().paragraph_format.space_after = Pt(30)
    icon = document.add_paragraph()
    icon.alignment = WD_ALIGN_PARAGRAPH.CENTER
    icon.add_run().add_picture(str(ROOT / "src/neomapper/presentation/data/neomapper_icon.png"), width=Inches(1.55))
    title = document.add_paragraph("Manual do Usuário do NEOMapper", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    remove_paragraph_borders(title)
    subtitle = document.add_paragraph("Consulta astronômica e planejamento observacional", style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    version = document.add_paragraph(f"Versão {APP_VERSION}   Windows   Setembro de 2026")
    version.alignment = WD_ALIGN_PARAGRAPH.CENTER
    version.paragraph_format.space_before = Pt(16)
    document.add_paragraph().paragraph_format.space_after = Pt(24)
    opening = document.add_paragraph()
    opening.alignment = WD_ALIGN_PARAGRAPH.CENTER
    opening.add_run(
        "Este manual acompanha o fluxo atual do NEOMapper: selecionar um objeto, consultar seus dados, "
        "avaliar altitude e eventos orbitais, configurar mapas, gerar animações e exportar relatórios."
    )
    document.add_paragraph().paragraph_format.space_after = Pt(18)
    warning = document.add_paragraph()
    warning.alignment = WD_ALIGN_PARAGRAPH.CENTER
    warning.add_run("Uso científico responsável. ").bold = True
    warning.add_run(
        "O NEOMapper auxilia planejamento e divulgação. Confirme observações críticas em fontes oficiais "
        "e não use o programa como único sistema para decisões de segurança."
    )
    page_break(document)

    heading(document, "Conteúdo e fluxo de trabalho")
    paragraph(document, "O programa organiza o trabalho em seis abas, sempre visíveis na coluna lateral.")
    table(document, ["Ordem", "Aba", "Finalidade"], [
        ["1", "Objeto", "Informar designação e data, pesquisar, usar favoritos e plotar o Sistema Solar."],
        ["2", "Info Obs", "Gerar gráfico de altitude, janela observacional e Resumo Orbital."],
        ["3", "Camadas", "Definir instante, tipo de mapa, camadas e abrir efemérides."],
        ["4", "Animação", "Definir intervalo, passo, velocidade e formato de saída."],
        ["5", "Ajustes", "Configurar local, horário, idioma, unidade, imagem e limites."],
        ["6", "Sobre", "Consultar versão, créditos, canal AstroNEOS e ajuda contextual."],
    ], [.65, 1.25, 5.1])
    heading(document, "Fluxo recomendado", 2)
    numbered(document, "Abra Ajustes e confirme o local, UTC ou LOCAL, idioma e unidade de distância.")
    numbered(document, "Em Objeto, digite a designação e a data de observação e clique em Buscar.")
    numbered(document, "Passe para Info Obs, gere o gráfico de altitude e, se desejar, o Resumo Orbital.")
    numbered(document, "Em Camadas, confirme data e hora, escolha o mapa e clique em Gerar mapa.")
    numbered(document, "Use Animação quando precisar acompanhar a evolução no tempo.")
    paragraph(document, "Os botões azuis indicam as ações principais. Durante consultas e renderizações, acompanhe a mensagem central e a barra de status inferior.")

    heading(document, "Instalação e primeiros ajustes")
    heading(document, "Instalador", 2)
    paragraph(document, "Baixe o arquivo NEOMapper-4.2.0-Windows-x64-Setup.exe na página oficial de Releases. Execute o instalador e abra o programa pelo menu Iniciar ou pelo atalho opcional da Área de Trabalho.")
    heading(document, "Edição portátil", 2)
    paragraph(document, "Extraia todo o conteúdo de NEOMapper-Windows-x64.zip para uma pasta e execute NEOMapper.exe. Não mova apenas o executável, pois ele depende da pasta interna que o acompanha.")
    heading(document, "Aviso do Windows", 2)
    paragraph(document, "A versão 4.2.0 não possui assinatura digital. O Windows pode mostrar um aviso de reputação. Confirme que o arquivo veio da página oficial e compare seu SHA 256 com SHA256SUMS.txt.")
    heading(document, "Antes da primeira busca", 2)
    bullet(document, "Abra Ajustes e revise o local de referência.")
    bullet(document, "Escolha UTC ou LOCAL; essa seleção afeta entradas e horários exibidos.")
    bullet(document, "Escolha UA ou km para todas as distâncias astronômicas.")
    bullet(document, "Mantenha HD para rapidez; use FHD, 2K ou 4K somente quando necessário.")
    paragraph(document, "As consultas precisam de internet. O NEOMapper tenta primeiro o JPL Horizons e usa o MPES do Minor Planet Center como contingência quando aplicável.")

    page_break(document)
    heading(document, "Aba Objeto")
    paragraph(document, "Esta é a entrada do fluxo. Informe uma designação reconhecida, como 99942, 1P ou 2026 RW3, selecione a data e clique em Buscar.")
    figure(document, "01_object.png", "Figura 1  Aba Objeto da versão 4.2.0")
    heading(document, "Resultado da busca", 2)
    paragraph(document, "Uma busca bem-sucedida mostra o nome resolvido, a distância ao Sol, a distância à Terra e a magnitude. Ela também habilita Plotar Sistema Solar e Continuar para Info Obs. Alterar o objeto ou a data invalida o resultado anterior e exige uma nova busca.")
    heading(document, "Plotar Sistema Solar", 2)
    paragraph(document, "O esquema usa a posição heliocêntrica e a órbita osculadora do objeto na data selecionada. Planetas e objeto aparecem na mesma projeção eclíptica. O desenho serve para contexto geométrico; não é uma integração numérica da órbita futura.")

    page_break(document)
    heading(document, "Favoritos")
    paragraph(document, "O botão Favoritos fica sempre disponível. A janela guarda designação, nome e notas no computador e não inicia consulta de rede.")
    figure(document, "07_favorites.png", "Figura 2  Cadastro e seleção de favoritos", width=5.25)
    bullet(document, "Para cadastrar, clique em Novo favorito, preencha os campos e escolha Salvar favorito.")
    bullet(document, "Para atualizar, selecione uma entrada, altere nome ou notas e salve novamente.")
    bullet(document, "Para preencher a aba Objeto, selecione a entrada e clique em Usar favorito.")
    bullet(document, "Para excluir, selecione a entrada e clique em Remover favorito.")
    paragraph(document, "Salvar novamente a mesma designação atualiza o registro existente. Novo favorito limpa o formulário para que outro objeto possa ser cadastrado.")

    page_break(document)
    heading(document, "Aba Info Obs")
    paragraph(document, "Esta aba reúne o planejamento da noite. O gráfico de altitude e o Resumo Orbital possuem botões independentes.")
    figure(document, "02_obs_information.png", "Figura 3  Gráfico de altitude e Resumo Orbital no painel central")
    heading(document, "Gráfico de altitude", 2)
    paragraph(document, "Plotar gráfico de altitude consulta 24 horas em amostras de cinco minutos. O gráfico mostra a altitude do objeto, faixas de dia e crepúsculo e a janela que atende aos limites escolhidos. O painel lateral recebe início, melhor horário, fim e altitude máxima a partir das mesmas amostras.")
    bullet(document, "Altitude mínima da janela define a elevação mínima aceitável para o objeto.")
    bullet(document, "Limite de altura do Sol define a escuridão exigida; menos 12 graus corresponde ao crepúsculo náutico.")
    bullet(document, "Mostrar Lua aparece abaixo do gráfico e alterna a curva lunar sem recalcular o objeto.")

    page_break(document)
    heading(document, "Resumo Orbital")
    paragraph(document, "Resumo orbital pode ser usado mesmo sem gerar o gráfico de altitude. Durante o cálculo, o programa mostra Calculando resumo orbital abaixo do gráfico ou no centro da área de prévia.")
    figure(document, "08_orbital_summary.png", "Figura 4  Resumo Orbital ampliado em unidades astronômicas", width=7.05)
    table(document, ["Linha", "Significado"], [
        ["Periélio", "Passagem mais próxima do Sol fornecida ou refinada para a solução orbital."],
        ["Data selecionada", "Efemérides no instante escolhido para o mapa."],
        ["Maior aproximação", "Menor distância à Terra encontrada no intervalo informado no quadro."],
    ], [1.65, 5.35])
    paragraph(document, "A tabela inclui magnitude, distâncias, ascensão reta, declinação, elongação e ângulo de fase. PA da cauda aparece somente para cometas. As distâncias seguem UA ou km definido em Ajustes.")
    heading(document, "Períodos longos", 2)
    paragraph(document, "Na maioria dos objetos, a maior aproximação é pesquisada em mais ou menos 365 dias da data selecionada. Para órbitas ligadas de período longo, como 1P Halley, o período osculador localiza a próxima passagem e o Horizons é consultado ao redor dela. O intervalo efetivamente pesquisado aparece sob a tabela. O resultado não representa a menor distância de toda a órbita.")
    bullet(document, "Ampliar resumo abre a mesma tabela em uma janela maior.")
    bullet(document, "Salvar PDF e Salvar CSV exportam os dados já calculados sem nova consulta.")
    bullet(document, "Ajuda abre explicações dos termos e das convenções científicas.")

    page_break(document)
    heading(document, "Aba Camadas")
    paragraph(document, "Camadas controla o instante e o conteúdo do mapa. A data vem inicialmente da aba Objeto, mas pode ser alterada aqui sem modificar a data da pesquisa original.")
    figure(document, "03_layers.png", "Figura 5  Camadas com o tipo Mapa celeste")
    heading(document, "Mapa de visibilidade", 2)
    paragraph(document, "Mostra em quais regiões o objeto está acima do horizonte e permite combinar Dia e Noite, crepúsculos civil, náutico e astronômico, terminador, ponto de referência, linhas de altitude e quadro da janela observacional.")
    heading(document, "Mapa celeste", 2)
    paragraph(document, "Mostra o céu do local configurado. O centro representa o zênite e a borda, o horizonte. A magnitude limite controla quantas estrelas aparecem. O objeto não é desenhado artificialmente no horizonte quando está abaixo de zero grau.")
    heading(document, "Gerar mapa", 2)
    paragraph(document, "Confirme data, hora e tipo e clique em Gerar mapa. A mensagem central informa a etapa atual. O PNG é gravado automaticamente na pasta do objeto e exibido na prévia.")

    page_break(document)
    heading(document, "Efemérides e relatório")
    paragraph(document, "O botão Efemérides fica na aba Camadas. Informe o intervalo e o passo solicitados na janela de efemérides. O limite máximo de linhas é definido em Ajustes e é verificado antes da consulta.")
    bullet(document, "Ascensão reta é exibida em horas, minutos e segundos.")
    bullet(document, "Declinação é exibida com sinal, graus, minutos e segundos.")
    bullet(document, "Movimento aparece em segundos de arco por minuto e o ângulo de posição é medido do norte para leste.")
    bullet(document, "Magnitude V ou magnitude total cometária T é identificada na tabela.")
    paragraph(document, "Ver relatório gera uma prévia PDF com todas as páginas. Na prévia, use Salvar PDF, Imprimir PDF ou Fechar. Salvar e imprimir usam o relatório já montado e não repetem a consulta.")
    paragraph(document, "Os relatórios de efemérides usam horários UTC. Confira objeto, intervalo, quantidade de linhas e procedência antes de arquivar ou imprimir.")

    page_break(document)
    heading(document, "Aba Animação")
    figure(document, "04_animation.png", "Figura 6  Controles atuais de animação")
    table(document, ["Controle", "Uso"], [
        ["Hora inicial e final", "Definem o intervalo no modo UTC ou LOCAL escolhido em Ajustes."],
        ["Passo", "Aceita minutos, hora, dia, semana, mês ou ano; passos menores geram mais quadros."],
        ["Velocidade", "Controla a reprodução e a taxa do GIF ou MP4, sem alterar as épocas calculadas."],
        ["Formato", "Gera GIF, MP4, ambos ou apenas quadros PNG."],
        ["Painel informativo", "Inclui dados do objeto no produto visual."],
        ["Rastro", "Mantém posições anteriores do objeto no mapa celeste."],
        ["Manter quadros", "Preserva os PNGs individuais após a animação."],
    ], [1.65, 5.35])
    paragraph(document, "Gerar animação valida previamente a quantidade de quadros. Cancelar interrompe uma geração em andamento. O progresso aparece no painel e na barra inferior. Abrir pasta das animações mostra o diretório de saída do objeto.")
    heading(document, "Player", 2)
    paragraph(document, "Após gerar, use Reproduzir, quadro anterior, quadro seguinte, controle deslizante e Repetir. GIFs exportados repetem continuamente. No mapa celeste, a sequência usa uma janela contínua acima do horizonte e termina quando o objeto deixa de ser visível.")

    page_break(document)
    heading(document, "Aba Ajustes")
    figure(document, "05_settings.png", "Figura 7  Local de referência e preferências reunidos em Ajustes")
    heading(document, "Local de referência", 2)
    paragraph(document, "O local define o horizonte topocêntrico, o fuso LOCAL e o ponto de referência. Escolha uma predefinição ou informe nome, latitude, longitude e altitude em metros.")
    table(document, ["Ação", "Resultado"], [
        ["Salvar", "Cria uma nova predefinição com os valores informados."],
        ["Atualizar", "Substitui os valores da predefinição selecionada."],
        ["Excluir", "Remove uma predefinição criada pelo usuário."],
        ["Personalizado", "Usa coordenadas sem depender de uma predefinição salva."],
    ], [1.35, 5.65])
    paragraph(document, "Latitude positiva indica norte; longitude positiva indica leste. Revise sinais e altitude antes de interpretar o horizonte ou o horário local.")

    page_break(document)
    heading(document, "Preferências e limites")
    table(document, ["Ajuste", "Efeito"], [
        ["Exibição do horário", "UTC ou LOCAL para entradas, títulos e datas apresentadas."],
        ["Idioma do mapa", "EN, PT ou ES para interface, mapas, ajuda e relatórios."],
        ["Tamanho da imagem", "HD, FHD, 2K ou 4K; tamanhos maiores exigem mais memória e tempo."],
        ["DPI", "Define a densidade do arquivo renderizado."],
        ["Unidade de distância", "UA ou km nos cards, mapas, Sistema Solar, resumo e animações novas."],
        ["Limite de quadros", "Interrompe pedidos de animação excessivos antes da renderização."],
        ["Limite de efemérides", "Interrompe pedidos com linhas demais antes da consulta."],
    ], [2.0, 5.0])
    paragraph(document, "As preferências são salvas automaticamente. Trocar a unidade atualiza cards e gráficos estáticos já carregados; animações antigas precisam ser geradas novamente para alterar seus textos.")
    heading(document, "UTC e LOCAL", 2)
    paragraph(document, "UTC é a escala de apresentação universal. LOCAL usa o fuso determinado pelas coordenadas do observador. Ao planejar uma observação próxima da meia noite, confira a data completa exibida, pois a conversão pode mudar o dia civil.")

    page_break(document)
    heading(document, "Aba Sobre e ajuda")
    figure(document, "06_about.png", "Figura 8  Sobre com versão, créditos e acesso à ajuda")
    paragraph(document, "Sobre apresenta a versão instalada, o propósito do software e os créditos do catálogo Hipparcos. O subconjunto estelar é creditado à ESA e está sujeito à licença CC BY NC 3.0 IGO.")
    paragraph(document, "O NEOMapper é um oferecimento do Observatório SONEAR e do canal AstroNEOS. O link de inscrição abre o canal no navegador padrão.")
    heading(document, "Ajuda contextual", 2)
    paragraph(document, "O botão Ajuda abre um glossário com ascensão reta, declinação, elongação, ângulo de fase, PA da cauda, magnitude, horários, referências do Resumo Orbital, limites de observação e unidades. Cabeçalhos e células científicas também possuem explicações ao passar o mouse.")
    paragraph(document, "A ajuda acompanha o idioma selecionado em Ajustes.")

    page_break(document)
    heading(document, "Arquivos e dados do usuário")
    paragraph(document, "O NEOMapper grava configurações, favoritos, logs e produtos fora da pasta instalada.")
    table(document, ["Conteúdo", "Local padrão"], [
        ["Configuração", r"%LOCALAPPDATA%\NEOMapper\NEOMapper_config.ini"],
        ["Favoritos", r"%LOCALAPPDATA%\NEOMapper\favorites.json"],
        ["Mapas e animações", r"%LOCALAPPDATA%\NEOMapper\output\<objeto>-<data>"],
        ["Logs", r"%LOCALAPPDATA%\NEOMapper\logs"],
    ], [1.65, 5.35])
    paragraph(document, "Cada pasta de saída combina uma designação segura para Windows com a data local de geração. Imagens, gráficos e animações do mesmo objeto e dia ficam reunidos.")
    heading(document, "Privacidade e rede", 2)
    paragraph(document, "O aplicativo não possui conta nem telemetria própria. Ao consultar um objeto, envia ao JPL Horizons ou ao Minor Planet Center a designação, o instante e os parâmetros necessários. Logs de erro ficam no computador e podem conter detalhes técnicos da consulta.")

    page_break(document)
    heading(document, "Solução de problemas")
    table(document, ["Sintoma", "Ação recomendada"], [
        ["Objeto não encontrado", "Revise a designação e a conexão. Aguarde e tente novamente se o provedor estiver indisponível."],
        ["Botões seguintes desabilitados", "Clique em Buscar novamente após alterar objeto ou data."],
        ["Gráfico sem janela verde", "Revise altitude mínima e limite solar; o objeto pode não atender aos dois critérios."],
        ["Resumo demora", "A busca pode consultar centenas de épocas e, em períodos longos, a próxima passagem orbital."],
        ["Maior aproximação não aparece", "A tabela compacta deve mostrar três linhas; maximize a janela e confirme que está usando a versão 4.2.0."],
        ["Mapa celeste sem objeto", "Confira a altitude. Abaixo do horizonte, o marcador não é deslocado artificialmente para a borda."],
        ["Animação recusada", "Aumente o passo, diminua o intervalo ou eleve conscientemente o limite de quadros em Ajustes."],
        ["MP4 falha", "Consulte o log da animação e tente GIF; o codificador pode estar indisponível."],
        ["Horário inesperado", "Confirme UTC ou LOCAL e revise latitude e longitude do observador."],
        ["Portátil não abre", "Extraia o ZIP inteiro e mantenha NEOMapper.exe junto da pasta interna."],
    ], [2.05, 4.95])

    page_break(document)
    heading(document, "Referência rápida")
    table(document, ["Objetivo", "Caminho"], [
        ["Consultar objeto", "Objeto → designação e data → Buscar"],
        ["Guardar objeto", "Objeto → Favoritos → Novo favorito → Salvar favorito"],
        ["Ver órbita", "Objeto → Buscar → Plotar Sistema Solar"],
        ["Planejar a noite", "Info Obs → limites → Plotar gráfico de altitude"],
        ["Calcular eventos", "Info Obs → Resumo orbital"],
        ["Exportar resumo", "Resumo calculado → Salvar PDF ou Salvar CSV"],
        ["Gerar mapa", "Camadas → data e hora → tipo → Gerar mapa"],
        ["Ver efemérides", "Camadas → Efemérides → Ver relatório"],
        ["Criar animação", "Animação → intervalo e passo → Gerar animação"],
        ["Alterar local", "Ajustes → Local de referência"],
        ["Alterar idioma ou unidade", "Ajustes → Idioma do mapa ou Unidade de distância"],
        ["Consultar termos", "Sobre → Ajuda ou Resumo Orbital → Ajuda"],
    ], [2.15, 4.85])
    heading(document, "Limitações", 2)
    bullet(document, "Não existe banco local completo de efemérides; a maioria das consultas depende da rede.")
    bullet(document, "Órbitas e previsões futuras podem mudar quando os provedores atualizam a solução orbital.")
    bullet(document, "A maior aproximação do resumo vale para o intervalo exibido, não para toda a história do objeto.")
    bullet(document, "O programa não controla telescópios e não publica produtos diretamente na web.")
    closing = document.add_paragraph()
    closing.alignment = WD_ALIGN_PARAGRAPH.CENTER
    closing.paragraph_format.space_before = Pt(16)
    closing.add_run(f"Fim do Manual do Usuário do NEOMapper {APP_VERSION}").bold = True

    document.core_properties.title = "Manual do Usuário do NEOMapper"
    document.core_properties.subject = f"Operação do NEOMapper versão {APP_VERSION}"
    document.core_properties.author = "NEOMapper"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)


if __name__ == "__main__":
    build_manual()
